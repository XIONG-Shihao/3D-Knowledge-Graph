import io
import json

import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter

from backend.config import Settings
from backend.main import create_app
from backend.ragflow import RAGFlowClient


@pytest.fixture
def client(tmp_path):
    app = create_app(Settings(data_dir=tmp_path, seed_demo=False))
    with TestClient(app) as client:
        yield client


def new_base(client, name="Research", **extra):
    response = client.post("/api/knowledge-bases", json={"name": name, **extra})
    assert response.status_code == 201, response.text
    return "/api/knowledge-bases/" + response.json()["id"]


def upload(client, base, text, name="notes.md"):
    blob = text.encode() if isinstance(text, str) else text
    response = client.post(base + "/documents", files={"file": (name, blob)})
    assert response.status_code == 202, response.text
    return response.json()["id"]


def test_upload_organize_search_graph_and_delete(client):
    base = new_base(client)
    doc_id = upload(
        client, base, "# Atlas\n\nAlice leads the research team.\n\n[[Alice]] --leads--> [[Atlas]]"
    )
    docs = client.get(base + "/documents").json()
    assert docs[0]["status"] == "ready"
    assert docs[0]["chunk_count"] > 0
    assert "Atlas" in docs[0]["topics"]
    result = client.post(base + "/search", json={"question": "Alice"}).json()
    assert result["results"][0]["doc_id"] == doc_id
    graph = client.get(base + "/graph").json()
    edge = next(e for e in graph["edges"] if e["relation"] == "leads")
    assert edge["source_documents"] == [doc_id]
    assert edge["source_chunks"]
    assert client.get(base + f"/documents/{doc_id}/download").content.startswith(b"# Atlas")
    assert client.delete(base + f"/documents/{doc_id}").status_code == 204
    assert client.get(base + "/graph").json()["nodes"] == []
    assert client.post(base + "/search", json={"question": "Alice"}).json()["results"] == []


def test_workspace_scope_is_enforced(client):
    first, second = new_base(client, "First"), new_base(client, "Second")
    doc = upload(client, first, "Private acquisition roadmap")
    assert client.get(second + f"/documents/{doc}").status_code == 404
    assert client.get(second + f"/documents/{doc}/download").status_code == 404
    assert client.delete(second + f"/documents/{doc}").status_code == 404
    assert client.post(second + "/search", json={"question": "acquisition"}).json()["results"] == []


def test_import_does_not_invent_cross_workspace_citations(client):
    first, second = new_base(client, "First"), new_base(client, "Second")
    doc = upload(client, first, "# Project\nAlice leads Atlas.")
    payload = {"nodes": [{"id": "a", "label": "Alice", "source_documents": [doc]}], "edges": []}
    response = client.post(
        second + "/graph/import", files={"file": ("graph.json", json.dumps(payload))}
    )
    assert response.status_code == 200
    assert response.json()["nodes"][0]["source_documents"] == []
    assert response.json()["nodes"][0]["source_refs"] == [doc]


def test_import_and_export_round_trip_then_reset(client):
    base = new_base(client)
    doc = upload(client, base, "# Evidence\nAlice works at Acme.")
    payload = {
        "nodes": [{"id": "a", "label": "Alice", "type": "person"}, {"id": "b", "label": "Acme"}],
        "edges": [
            {"source": "a", "target": "b", "relation": "works at", "source_documents": [doc]}
        ],
    }
    assert (
        client.post(
            base + "/graph/import", files={"file": ("g.json", json.dumps(payload))}
        ).status_code
        == 200
    )
    exported = client.get(base + "/graph").json()
    assert exported["edges"][0]["source_documents"] == [doc]
    round_trip = client.post(
        base + "/graph/import", files={"file": ("g.json", json.dumps(exported))}
    )
    assert round_trip.status_code == 200
    assert round_trip.json()["nodes"][0]["label"] == "Alice"
    assert client.delete(base + "/graph/import").status_code == 204
    assert client.get(base + "/graph").json()["origin"] == "local"


@pytest.mark.parametrize(
    "payload",
    [
        [],
        {},
        {"nodes": "bad"},
        {"nodes": [{}]},
        {"nodes": [{"id": "a"}], "edges": [{"source": "a", "target": "missing"}]},
    ],
)
def test_bad_graphs_are_rejected(client, payload):
    base = new_base(client)
    response = client.post(base + "/graph/import", files={"file": ("g.json", json.dumps(payload))})
    assert response.status_code == 422


def test_failed_document_is_visible_and_can_be_retried(client):
    base = new_base(client)
    doc = upload(client, base, b"broken pdf", "broken.pdf")
    detail = client.get(base + f"/documents/{doc}").json()
    assert detail["status"] == "failed"
    assert detail["error"]
    assert client.post(base + f"/documents/{doc}/retry").status_code == 202
    assert client.get(base + f"/documents/{doc}").json()["status"] == "failed"


def test_size_type_empty_validation_and_safe_filename(client):
    base = new_base(client)
    assert client.post(base + "/documents", files={"file": ("a.exe", b"abc")}).status_code == 415
    assert client.post(base + "/documents", files={"file": ("a.txt", b"")}).status_code == 422
    client.app.state.settings.max_upload_bytes = 10
    assert client.post(base + "/documents", files={"file": ("a.txt", b"x" * 11)}).status_code == 413
    doc = upload(client, base, "safe", "../../secret.txt")
    assert client.get(base + f"/documents/{doc}").json()["name"] == "secret.txt"


def test_docx_parsing_includes_tables(client):
    doc = Document()
    doc.add_paragraph("The project leader is Alice.")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Marcus maintains RAGFlow."
    stream = io.BytesIO()
    doc.save(stream)
    base = new_base(client)
    uploaded = upload(client, base, stream.getvalue(), "team.docx")
    detail = client.get(base + f"/documents/{uploaded}").json()
    assert detail["status"] == "ready"
    assert "Marcus" in detail["chunks"][0]["content"]


def test_scanned_pdf_is_explicit_failure(client):
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    stream = io.BytesIO()
    writer.write(stream)
    base = new_base(client)
    uploaded = upload(client, base, stream.getvalue(), "scan.pdf")
    detail = client.get(base + f"/documents/{uploaded}").json()
    assert detail["status"] == "failed"
    assert "OCR" in detail["error"]


def test_plain_text_does_not_create_inferred_relationships(client):
    base = new_base(client)
    upload(client, base, "Alice and Atlas appear together in this document.")
    graph = client.get(base + "/graph").json()
    assert all(e["relation"] == "covers" for e in graph["edges"])


def test_chinese_keyword_retrieval(client):
    base = new_base(client)
    upload(client, base, "# 项目计划\n张伟负责知识图谱开发和文档检索。")
    assert client.post(base + "/search", json={"question": "知识图谱"}).json()["results"]


@pytest.mark.parametrize("encoding", ["utf-8", "utf-16", "gb18030"])
def test_chinese_encodings_and_relationships(client, encoding):
    base = new_base(client, "中文知识库")
    text = "# 知识图谱项目\n张伟负责文档检索。\n[[张伟]] --负责--> [[知识图谱项目]]"
    doc = upload(client, base, text.encode(encoding), "项目说明.txt")
    detail = client.get(base + f"/documents/{doc}").json()
    assert detail["status"] == "ready"
    assert "张伟负责文档检索" in detail["chunks"][0]["content"]
    assert client.post(base + "/search", json={"question": "知识图谱"}).json()["results"]
    assert "知识图谱项目" in detail["topics"]
    graph = client.get(base + "/graph").json()
    edge = next(e for e in graph["edges"] if e["relation"] == "负责")
    assert edge["source_documents"] == [doc]


def test_persistence_across_restart(tmp_path):
    settings = Settings(data_dir=tmp_path, seed_demo=False)
    with TestClient(create_app(settings)) as client:
        base = new_base(client)
        upload(client, base, "Persistent knowledge")
    with TestClient(create_app(settings)) as client:
        assert len(client.get(base + "/documents").json()) == 1
        assert client.post(base + "/search", json={"question": "Persistent"}).json()["results"]


def test_demo_seed_is_repeatable(tmp_path):
    settings = Settings(data_dir=tmp_path, seed_demo=True)
    for _ in range(2):
        with TestClient(create_app(settings)) as client:
            bases = client.get("/api/knowledge-bases").json()
            assert len(bases) == 1
            assert bases[0]["document_count"] == 5


def test_ragflow_end_to_end_http_contract(tmp_path):
    requests = []

    def server(request):
        requests.append(request)
        assert request.headers["Authorization"] == "Bearer test-key"
        path = request.url.path
        if path == "/api/v1/datasets" and request.method == "POST":
            assert json.loads(request.content)["permission"] == "me"
            data = {"id": "remote-base"}
        elif path.endswith("/documents") and request.method == "POST":
            assert b"team.md" in request.content
            data = [{"id": "remote-doc"}]
        elif path.endswith("/documents") and request.method == "GET":
            data = {"docs": [{"id": "remote-doc", "name": "team.md", "run": "DONE", "progress": 1}]}
        elif path.endswith("/documents/remote-doc/chunks"):
            data = {
                "chunks": [
                    {
                        "id": "remote-chunk",
                        "content": "<p>Alice leads Atlas.</p>",
                        "important_keywords": ["Atlas"],
                    }
                ]
            }
        elif path == "/api/v1/retrieval":
            assert json.loads(request.content)["dataset_ids"] == ["remote-base"]
            data = {
                "chunks": [
                    {
                        "id": "remote-chunk",
                        "document_id": "remote-doc",
                        "content": "Alice leads Atlas.",
                        "similarity": 0.9,
                    }
                ]
            }
        elif path.endswith("/knowledge_graph"):
            data = {
                "graph": {
                    "nodes": [
                        {
                            "id": "Alice",
                            "entity_name": "Alice",
                            "entity_type": "PERSON",
                            "source_id": ["remote-chunk"],
                        },
                        {"id": "Atlas", "entity_name": "Atlas"},
                    ],
                    "edges": [
                        {
                            "source": "Alice",
                            "target": "Atlas",
                            "description": "Alice leads Atlas.",
                            "source_id": ["remote-chunk"],
                        }
                    ],
                }
            }
        else:
            data = True
        return httpx.Response(200, json={"code": 0, "data": data})

    settings = Settings(
        data_dir=tmp_path,
        seed_demo=False,
        ragflow_url="http://ragflow.test",
        ragflow_key="test-key",
        legacy_graph_api=True,
    )
    rag = RAGFlowClient(settings, transport=httpx.MockTransport(server))
    with TestClient(create_app(settings, rag)) as client:
        base = new_base(client, backend="ragflow")
        doc = upload(client, base, "Alice leads Atlas.", "team.md")
        assert client.get(base + f"/documents/{doc}").json()["status"] == "processing"
        assert client.post(base + "/sync").status_code == 200
        detail = client.get(base + f"/documents/{doc}").json()
        assert detail["status"] == "ready"
        assert detail["chunks"][0]["content"] == "Alice leads Atlas."
        assert (
            client.post(base + "/search", json={"question": "Alice"}).json()["results"][0]["doc_id"]
            == doc
        )
        assert client.post(base + "/graph/sync").json()["status"] == "ready"
        graph = client.get(base + "/graph").json()
        assert graph["nodes"][0]["source_documents"] == [doc]
        assert graph["edges"][0]["source_documents"] == [doc]
        assert client.delete(base + f"/documents/{doc}").status_code == 204
        assert client.get(base + "/graph").json()["origin"] == "local"
    assert any(
        r.method == "DELETE" and json.loads(r.content) == {"ids": ["remote-doc"]} for r in requests
    )


def test_ragflow_failures_are_visible_without_leaking_key(tmp_path):
    settings = Settings(
        data_dir=tmp_path,
        seed_demo=False,
        ragflow_url="http://ragflow.test",
        ragflow_key="secret-key",
    )
    rag = RAGFlowClient(settings, transport=httpx.MockTransport(lambda _: httpx.Response(401)))
    with TestClient(create_app(settings, rag)) as client:
        response = client.post(
            "/api/knowledge-bases", json={"name": "remote", "backend": "ragflow"}
        )
        assert response.status_code == 502
        assert "secret-key" not in response.text
        assert "secret-key" not in client.get("/api/config").text


def test_remote_cancel_does_not_leave_searchable_chunks(tmp_path):
    run = "DONE"

    def server(request):
        if request.url.path.endswith("/documents"):
            data = {"docs": [{"id": "d", "name": "d.md", "run": run, "progress": 1}]}
        elif request.url.path.endswith("/chunks"):
            data = {"chunks": [{"id": "c", "content": "Stale evidence"}]}
        else:
            data = [{"id": "k"}]
        return httpx.Response(200, json={"code": 0, "data": data})

    settings = Settings(
        data_dir=tmp_path, seed_demo=False, ragflow_url="http://ragflow.test", ragflow_key="key"
    )
    with TestClient(
        create_app(settings, RAGFlowClient(settings, httpx.MockTransport(server)))
    ) as client:
        base = new_base(client, backend="ragflow", dataset_id="k")
        client.post(base + "/sync")
        assert client.app.state.store.one("SELECT COUNT(*) AS n FROM chunks")["n"] == 1
        run = "CANCEL"
        client.post(base + "/sync")
        assert client.app.state.store.one("SELECT COUNT(*) AS n FROM chunks")["n"] == 0
        assert client.get(base + "/documents").json()[0]["status"] == "failed"
