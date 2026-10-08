import hashlib
import io
import json
import re
import zipfile
from collections import Counter
from html import unescape
from html.parser import HTMLParser
from pathlib import Path

SUPPORTED = {".txt", ".md", ".pdf", ".docx"}
STOP_WORDS = set(
    """the and for with from this that are was were has have will can our your
into how what when where which their they not but its you all use using about also each
through than then these those a an of to in on is it as be by or at we us may should
must does do used document project notes overview introduction who why""".split()
)


class TextHTMLParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts = []

    def handle_data(self, data):
        self.parts.append(data)

    def handle_starttag(self, tag, attrs):
        if tag in {"p", "br", "div", "tr", "li", "h1", "h2", "h3"}:
            self.parts.append("\n")


def plain_text(text):
    if re.search(r"</?(?:p|div|table|span|em|strong|a|h[1-6])\b", text, re.I):
        parser = TextHTMLParser()
        parser.feed(text)
        return unescape("".join(parser.parts)).strip()
    return text


def parse_file(name, blob):
    ext = Path(name).suffix.lower()
    if ext not in SUPPORTED:
        raise ValueError("Supported formats: Markdown, TXT, PDF, DOCX.")
    if ext in {".md", ".txt"}:
        try:
            pages = [(None, blob.decode("utf-8-sig"))]
        except UnicodeDecodeError as exc:
            try:
                encoding = "utf-16" if blob.startswith((b"\xff\xfe", b"\xfe\xff")) else "gb18030"
                pages = [(None, blob.decode(encoding))]
            except UnicodeDecodeError:
                raise ValueError("Use UTF-8, UTF-16 with BOM, or GB18030 for text files.") from exc
    elif ext == ".pdf":
        from pypdf import PdfReader

        try:
            reader = PdfReader(io.BytesIO(blob))
            if reader.is_encrypted:
                raise ValueError("Encrypted PDFs are unsupported; upload an unlocked copy.")
            if len(reader.pages) > 1000:
                raise ValueError("This MVP supports PDFs with up to 1,000 pages.")
            pages = [(i + 1, page.extract_text() or "") for i, page in enumerate(reader.pages)]
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("Could not read this PDF. Check that the file is valid.") from exc
    else:
        from docx import Document

        try:
            with zipfile.ZipFile(io.BytesIO(blob)) as archive:
                if sum(x.file_size for x in archive.infolist()) > 50 * 1024 * 1024:
                    raise ValueError("The expanded DOCX exceeds the 50 MB limit.")
            doc = Document(io.BytesIO(blob))
            text = "\n\n".join(p.text for p in doc.paragraphs)
            text += "\n" + "\n".join(
                " | ".join(c.text for c in row.cells) for t in doc.tables for row in t.rows
            )
            pages = [(None, text)]
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("Could not read this DOCX. Check that the file is valid.") from exc
    text = "\n\n".join(t for _, t in pages)
    if not text.strip():
        raise ValueError("No readable text. Scanned PDFs need OCR through RAGFlow.")
    if len(text) > 2_000_000:
        raise ValueError("Extracted text exceeds the MVP limit of 2 million characters.")
    chunks = []
    for page, content in pages:
        content = content.strip()
        start = 0
        while start < len(content):
            end = min(start + 1000, len(content))
            if end < len(content):
                split = content.rfind("\n", start + 500, end)
                if split < 0:
                    split = content.rfind(" ", start + 500, end)
                if split > start:
                    end = split
            part = content[start:end].strip()
            if part:
                chunks.append({"content": part, "page": page})
            if end >= len(content):
                break
            start = max(start + 1, end - 100)
    return chunks, extract_topics(text)


def tokenize(text):
    # Character bigrams make Chinese documents usable in local keyword search too.
    words = re.findall(r"[a-zA-Z][a-zA-Z0-9_-]{2,}", text.lower())
    for phrase in re.findall(r"[\u4e00-\u9fff]+", text):
        words.extend(phrase[i : i + 2] for i in range(max(1, len(phrase) - 1)))
    return [w for w in words if w not in STOP_WORDS]


def extract_topics(text):
    headings = re.findall(r"^#{1,6}\s+(.+)$", text, re.M)
    links = re.findall(r"\[\[([^\]\n]{1,100})\]\]", text)
    named = list(dict.fromkeys(h.strip() for h in headings + links if h.strip()))[:8]
    if named:
        return named
    return [word.title() for word, _ in Counter(tokenize(text)).most_common(5)]


def stable_id(prefix, label):
    return prefix + hashlib.sha256(label.casefold().strip().encode()).hexdigest()[:16]


def local_graph(documents, chunks):
    nodes, edges = {}, {}

    def add_node(node):
        if node["id"] in nodes:
            old = nodes[node["id"]]
            old["source_documents"] = sorted(
                set(old["source_documents"] + node["source_documents"])
            )
        else:
            nodes[node["id"]] = node

    for doc in documents:
        if doc["status"] != "ready":
            continue
        doc_node = "doc:" + doc["id"]
        add_node(
            {
                "id": doc_node,
                "label": doc["name"],
                "type": "document",
                "description": "Original source document.",
                "source_documents": [doc["id"]],
                "origin": "local",
            }
        )
        for topic in json.loads(doc["topics"]):
            topic_id = stable_id("topic:", topic)
            add_node(
                {
                    "id": topic_id,
                    "label": topic,
                    "type": "topic",
                    "description": "A heading, explicit reference, or frequent keyword in a document.",
                    "source_documents": [doc["id"]],
                    "origin": "local",
                }
            )
            key = stable_id("edge:", doc_node + topic_id)
            edges[key] = {
                "id": key,
                "source": doc_node,
                "target": topic_id,
                "relation": "covers",
                "description": "Document/topic connection.",
                "source_documents": [doc["id"]],
                "origin": "local",
            }
    for chunk in chunks:
        pattern = r"\[\[([^\]\n]+)\]\]\s*--([\w -]+)-->\s*\[\[([^\]\n]+)\]\]"
        for match in re.finditer(pattern, chunk["content"]):
            source_label, relation, target_label = [v.strip() for v in match.groups()]
            source, target = [stable_id("topic:", v) for v in (source_label, target_label)]
            for label, node_id in [(source_label, source), (target_label, target)]:
                add_node(
                    {
                        "id": node_id,
                        "label": label,
                        "type": "topic",
                        "description": "Explicitly annotated in a source document.",
                        "source_documents": [chunk["doc_id"]],
                        "origin": "local",
                    }
                )
            key = stable_id("edge:", source + relation + target)
            if key not in edges:
                edges[key] = {
                    "id": key,
                    "source": source,
                    "target": target,
                    "relation": relation,
                    "description": match.group(0),
                    "source_documents": [],
                    "source_chunks": [],
                    "origin": "annotation",
                }
            edge = edges[key]
            edge["source_documents"] = sorted(set(edge["source_documents"] + [chunk["doc_id"]]))
            edge["source_chunks"] = sorted(set(edge["source_chunks"] + [chunk["id"]]))
    return {"nodes": list(nodes.values()), "edges": list(edges.values()), "origin": "local"}


def normalize_graph(payload, documents, chunks, origin="imported"):
    if not isinstance(payload, dict):
        raise ValueError("Graph JSON must be an object with nodes and edges.")
    graph = payload.get("data", payload)
    if not isinstance(graph, dict):
        raise ValueError("Invalid graph data.")
    graph = graph.get("graph", graph)
    if not isinstance(graph, dict):
        raise ValueError("Invalid graph object.")
    if "nodes" not in graph:
        raise ValueError("Graph JSON must contain a nodes array.")
    raw_nodes = graph.get("nodes", [])
    raw_edges = graph.get("edges", graph.get("links", []))
    if not isinstance(raw_nodes, list) or not isinstance(raw_edges, list):
        raise ValueError("nodes and edges must be arrays.")
    if len(raw_nodes) > 2000 or len(raw_edges) > 5000:
        raise ValueError("This MVP supports up to 2,000 nodes and 5,000 edges per imported graph.")
    doc_map = {d["id"]: d["id"] for d in documents}
    doc_map.update({d["remote_id"]: d["id"] for d in documents if d.get("remote_id")})
    chunk_map = {c["remote_id"]: c["doc_id"] for c in chunks if c.get("remote_id")}
    chunk_map.update({c["id"]: c["doc_id"] for c in chunks})

    def sources(item):
        refs = item.get("source_documents", item.get("source_id", [])) or []
        if isinstance(refs, str):
            refs = refs.split("<SEP>")
        if not isinstance(refs, list):
            raise ValueError("Source references must be an array or a string.")
        supplied = item.get("source_refs", []) or []
        if not isinstance(supplied, list):
            raise ValueError("source_refs must be an array.")
        refs = list(dict.fromkeys(str(ref) for ref in refs + supplied))
        return sorted(
            {doc_map.get(r) or chunk_map[r] for r in refs if r in doc_map or r in chunk_map}
        ), refs

    nodes, aliases = {}, {}
    for item in raw_nodes:
        if not isinstance(item, dict):
            raise ValueError("Each node must be an object.")
        label = str(item.get("label") or item.get("entity_name") or item.get("id") or "").strip()
        if not label or len(label) > 300:
            raise ValueError("Every node needs a non-empty label of up to 300 characters.")
        external_id = str(item.get("id") or label)
        node_id = stable_id("entity:", external_id)
        if node_id in nodes:
            raise ValueError("Duplicate node IDs in graph.")
        aliases[external_id] = node_id
        aliases[label] = node_id
        docs, refs = sources(item)
        nodes[node_id] = {
            "id": node_id,
            "label": label,
            "type": str(item.get("type") or item.get("entity_type") or "concept").lower(),
            "description": str(item.get("description") or ""),
            "source_documents": docs,
            "source_refs": refs,
            "origin": origin,
        }
    edges = []
    for i, item in enumerate(raw_edges):
        if not isinstance(item, dict):
            raise ValueError("Each edge must be an object.")
        source = aliases.get(str(item.get("source", "")))
        target = aliases.get(str(item.get("target", "")))
        if not source or not target:
            raise ValueError("Every edge must reference existing nodes.")
        docs, refs = sources(item)
        relation = str(item.get("relation") or item.get("label") or "related to")
        edges.append(
            {
                "id": f"remote-edge:{i}",
                "source": source,
                "target": target,
                "relation": relation,
                "description": str(item.get("description") or ""),
                "source_documents": docs,
                "source_refs": refs,
                "origin": origin,
            }
        )
    return {"nodes": list(nodes.values()), "edges": edges, "origin": origin}


def search_chunks(query, chunks, limit=12):
    terms = Counter(tokenize(query))
    if not terms:
        return []
    corpus = [Counter(tokenize(c["content"])) for c in chunks]
    document_frequency = Counter(term for words in corpus for term in words)
    import math

    avg_length = sum(sum(words.values()) for words in corpus) / max(1, len(corpus))
    results = []
    for chunk, words in zip(chunks, corpus, strict=True):
        length = sum(words.values())
        score = 0.0
        for term in terms:
            freq = words[term]
            if freq:
                idf = math.log(
                    1
                    + (len(corpus) - document_frequency[term] + 0.5)
                    / (document_frequency[term] + 0.5)
                )
                score += (
                    idf * freq * 2.2 / (freq + 1.2 * (0.25 + 0.75 * length / max(1, avg_length)))
                )
        if score:
            results.append({**chunk, "score": round(score, 4)})
    return sorted(results, key=lambda c: c["score"], reverse=True)[:limit]
